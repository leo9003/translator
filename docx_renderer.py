from pathlib import Path
from typing import Optional

from docx import Document
from docx.document import Document as DocumentObject
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from docx.oxml.text.paragraph import CT_P
from docx.oxml.table import CT_Tbl

from docx.shared import Pt

from models import (
    Block,
    DocumentIR,
    ParagraphBlock,
    RunIR,
    TableBlock,
)


class DocxRenderer:
    """
    Template-based DOCX renderer.

    核心概念：

    不建立新的空白 Word 文件，而是直接打開原始 DOCX，
    只更新需要翻譯的 paragraph 文字。

    因此可以保留原始文件中的：

    - numbering.xml
    - styles.xml
    - theme
    - section settings
    - page settings
    - table properties
    - merged cells
    - borders / shading
    - relationships
    - list definitions
    - paragraph properties
    - 其他 OOXML metadata

    Parser 與 Renderer 必須使用相同 traversal 順序，
    讓：

        p_0000
        p_0001
        p_0002
        ...

    可以穩定對應回原始 DOCX paragraph。
    """

    # =========================================================
    # Public API
    # =========================================================

    def render(
        self,
        source_path: str,
        document_ir: DocumentIR,
        output_path: str,
    ) -> None:

        source_file = Path(source_path)
        output_file = Path(output_path)

        # -----------------------------------------------------
        # Validate source
        # -----------------------------------------------------

        if not source_file.exists():
            raise FileNotFoundError(
                f"Source DOCX not found: {source_file}"
            )

        if source_file.suffix.lower() != ".docx":
            raise ValueError(
                "Template renderer only supports .docx files."
            )

        # -----------------------------------------------------
        # IMPORTANT
        #
        # 使用原始 DOCX 當 template。
        #
        # 不再：
        #
        #     document = Document()
        #
        # -----------------------------------------------------

        document = Document(
            str(source_file)
        )

        # -----------------------------------------------------
        # 建立：
        #
        # paragraph_id -> ParagraphBlock
        #
        # -----------------------------------------------------

        paragraph_map: dict[
            str,
            ParagraphBlock,
        ] = {}

        self._collect_ir_paragraphs(
            blocks=document_ir.blocks,
            paragraph_map=paragraph_map,
        )

        # -----------------------------------------------------
        # paragraph_index 使用 list 是為了 recursive traversal
        # 時可以共享 mutable counter。
        # -----------------------------------------------------

        paragraph_index = [0]

        # -----------------------------------------------------
        # Walk original DOCX
        # -----------------------------------------------------

        self._update_container(
            container=document,
            paragraph_map=paragraph_map,
            paragraph_index=paragraph_index,
        )

        # -----------------------------------------------------
        # Validation
        #
        # 如果兩邊 paragraph 數量不同，通常代表：
        #
        # 1. source_path 不是產生這份 DocumentIR 的原始檔
        # 2. Parser / Renderer traversal 不一致
        # -----------------------------------------------------

        visited_count = paragraph_index[0]
        expected_count = len(paragraph_map)

        if visited_count != expected_count:
            raise ValueError(
                "Template rendering paragraph count mismatch.\n"
                f"Template paragraphs: {visited_count}\n"
                f"DocumentIR paragraphs: {expected_count}\n"
                "The source DOCX and DocumentIR may not come "
                "from the same document."
            )

        # -----------------------------------------------------
        # Save
        # -----------------------------------------------------

        output_file.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        document.save(
            str(output_file)
        )

    # =========================================================
    # Collect ParagraphBlock from DocumentIR
    # =========================================================

    def _collect_ir_paragraphs(
        self,
        blocks: list[Block],
        paragraph_map: dict[str, ParagraphBlock],
    ) -> None:

        for block in blocks:

            # -------------------------------------------------
            # Paragraph
            # -------------------------------------------------

            if isinstance(
                block,
                ParagraphBlock,
            ):

                if block.id in paragraph_map:
                    raise ValueError(
                        "Duplicate ParagraphBlock ID: "
                        f"{block.id}"
                    )

                paragraph_map[
                    block.id
                ] = block

                continue

            # -------------------------------------------------
            # Table
            # -------------------------------------------------

            if isinstance(
                block,
                TableBlock,
            ):

                for row in block.rows:

                    for cell in row.cells:

                        self._collect_ir_paragraphs(
                            blocks=cell.blocks,
                            paragraph_map=paragraph_map,
                        )

                continue

            raise TypeError(
                "Unsupported block type: "
                f"{type(block)}"
            )

    # =========================================================
    # Traverse Original DOCX
    # =========================================================

    def _update_container(
        self,
        container: DocumentObject | _Cell,
        paragraph_map: dict[str, ParagraphBlock],
        paragraph_index: list[int],
    ) -> None:

        # -----------------------------------------------------
        # 這裡刻意與 DocxParser._parse_container()
        # 使用相同邏輯。
        # -----------------------------------------------------

        if isinstance(
            container,
            DocumentObject,
        ):

            parent_element = (
                container.element.body
            )

        elif isinstance(
            container,
            _Cell,
        ):

            parent_element = (
                container._tc
            )

        else:
            raise TypeError(
                "Unsupported container type: "
                f"{type(container)}"
            )

        # -----------------------------------------------------
        # 按照原始 XML order traversal
        # -----------------------------------------------------

        for child in (
            parent_element.iterchildren()
        ):

            # -------------------------------------------------
            # Paragraph
            # -------------------------------------------------

            if isinstance(
                child,
                CT_P,
            ):

                paragraph = Paragraph(
                    child,
                    container,
                )

                paragraph_id = (
                    f"p_{paragraph_index[0]:04d}"
                )

                paragraph_index[0] += 1

                block = paragraph_map.get(
                    paragraph_id
                )

                if block is None:
                    raise ValueError(
                        "Cannot find ParagraphBlock "
                        f"{paragraph_id} "
                        "while rendering template."
                    )

                self._update_paragraph(
                    paragraph=paragraph,
                    block=block,
                )

            # -------------------------------------------------
            # Table
            # -------------------------------------------------

            elif isinstance(
                child,
                CT_Tbl,
            ):

                table = Table(
                    child,
                    container,
                )

                self._update_table(
                    table=table,
                    paragraph_map=paragraph_map,
                    paragraph_index=paragraph_index,
                )

    # =========================================================
    # Update Table
    # =========================================================

    def _update_table(
        self,
        table: Table,
        paragraph_map: dict[str, ParagraphBlock],
        paragraph_index: list[int],
    ) -> None:

        """
        不重建 Table。

        直接 traversal 原始 Table，因此可以保留：

        - column widths
        - row heights
        - borders
        - shading
        - table style
        - merged cells
        - cell properties
        - table layout
        """

        for row in table.rows:

            for cell in row.cells:

                self._update_container(
                    container=cell,
                    paragraph_map=paragraph_map,
                    paragraph_index=paragraph_index,
                )

    # =========================================================
    # Update Paragraph
    # =========================================================

    def _update_paragraph(
        self,
        paragraph: Paragraph,
        block: ParagraphBlock,
    ) -> None:

        original_text = paragraph.text
        translated_text = block.text

        # -----------------------------------------------------
        # 最重要的 Template Renderer 規則：
        #
        # 文字沒有改變 → 完全不要碰原始 XML。
        #
        # preserve 的內容因此可以 100% 保留原始 runs。
        #
        # 例如：
        #
        # COUNTYNAME
        # year
        # crop
        # CDD
        # DOI
        # Reference
        #
        # -----------------------------------------------------

        if original_text == translated_text:
            return

        # -----------------------------------------------------
        # 只有真的翻譯過的 paragraph 才修改。
        # -----------------------------------------------------

        self._replace_paragraph_text(
            paragraph=paragraph,
            translated_text=translated_text,
            translated_runs=block.runs,
        )

    # =========================================================
    # Replace Paragraph Text
    # =========================================================

    def _replace_paragraph_text(
        self,
        paragraph: Paragraph,
        translated_text: str,
        translated_runs: list[RunIR],
    ) -> None:

        """
        v1 translation replacement strategy.

        保留原始 paragraph 本身，因此：

            <w:p>
                <w:pPr>
                    numbering
                    style
                    spacing
                    indentation
                    pagination
                    ...
                </w:pPr>

                <w:r>...</w:r>
            </w:p>

        pPr 不會被重建。

        目前 translated paragraph 會集中到一個主要 run。

        未來 Step 10 才處理：
        semantic span -> multiple translated runs。
        """

        # -----------------------------------------------------
        # IMPORTANT BUG FIX
        #
        # paragraph.runs 每次存取都可能建立新的 Python
        # Run wrapper。
        #
        # 所以不能：
        #
        #     run is template_run
        #
        # 來判斷是不是同一個 Word run。
        #
        # 這裡一次取得 runs，後續全部使用同一份 list。
        # -----------------------------------------------------

        runs = list(
            paragraph.runs
        )

        # -----------------------------------------------------
        # Paragraph 沒有任何 run
        # -----------------------------------------------------

        if not runs:

            new_run = paragraph.add_run(
                translated_text
            )

            self._apply_ir_run_format(
                run=new_run,
                translated_runs=translated_runs,
            )

            return

        # -----------------------------------------------------
        # 找第一個真正有文字的 run。
        #
        # 預設使用 index 0。
        # -----------------------------------------------------

        template_index = 0

        for index, run in enumerate(
            runs
        ):

            if run.text.strip():

                template_index = index

                break

        # -----------------------------------------------------
        # 保存同一個 Run wrapper。
        # -----------------------------------------------------

        template_run = runs[
            template_index
        ]

        # -----------------------------------------------------
        # Step 1
        #
        # 清除 paragraph 中所有原始 visible text。
        #
        # 注意：
        #
        # run 本身沒有被刪掉。
        #
        # rPr 等 formatting metadata 仍然存在。
        # -----------------------------------------------------

        for run in runs:

            run.text = ""

        # -----------------------------------------------------
        # Step 2
        #
        # 把翻譯後的完整 paragraph 放到 template run。
        #
        # 這裡一定要在 clear 之後執行。
        # -----------------------------------------------------

        template_run.text = (
            translated_text
        )

        # -----------------------------------------------------
        # Step 3
        #
        # Mapper v1 通常會建立一個 translated RunIR，
        # 使用第一個 source run 的格式。
        #
        # 把已知格式套回 template run。
        # -----------------------------------------------------

        self._apply_ir_run_format(
            run=template_run,
            translated_runs=translated_runs,
        )

    # =========================================================
    # Apply Translation Run Formatting
    # =========================================================

    def _apply_ir_run_format(
        self,
        run,
        translated_runs: list[RunIR],
    ) -> None:

        """
        將 TranslationMapper 產生的 RunIR formatting
        套到 translated run。

        None 的 property 不覆蓋 template 原本的設定。
        """

        if not translated_runs:
            return

        # -----------------------------------------------------
        # 找第一個有內容的 translated RunIR。
        # -----------------------------------------------------

        run_ir: Optional[
            RunIR
        ] = None

        for candidate in (
            translated_runs
        ):

            if candidate.text.strip():

                run_ir = candidate

                break

        # -----------------------------------------------------
        # fallback
        # -----------------------------------------------------

        if run_ir is None:

            run_ir = (
                translated_runs[0]
            )

        # -----------------------------------------------------
        # Bold
        # -----------------------------------------------------

        if run_ir.bold is not None:

            run.bold = (
                run_ir.bold
            )

        # -----------------------------------------------------
        # Italic
        # -----------------------------------------------------

        if run_ir.italic is not None:

            run.italic = (
                run_ir.italic
            )

        # -----------------------------------------------------
        # Underline
        # -----------------------------------------------------

        if run_ir.underline is not None:

            run.underline = (
                run_ir.underline
            )

        # -----------------------------------------------------
        # Font
        # -----------------------------------------------------

        if run_ir.font_name:

            run.font.name = (
                run_ir.font_name
            )

        # -----------------------------------------------------
        # Font size
        # -----------------------------------------------------

        if (
            run_ir.font_size
            is not None
        ):

            run.font.size = Pt(
                run_ir.font_size
            )